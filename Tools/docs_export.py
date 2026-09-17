"""
docs_export.py — export a Docs/ Markdown file to a compact, shareable Word file.

Pandoc's default Word styles are very loose (big paragraph gaps, unbordered tables).
This converts with pandoc, then tightens the styles with python-docx.

Usage:
    python3 Tools/docs_export.py Docs/README.md ~/Desktop/VRDots_share/overview.docx

Requires: pandoc, python-docx (pip install python-docx).
"""

import re
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Helvetica Neue"
MONO = "Menlo"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
REPO_DOCS = "https://github.com/GeneStoner/UnityQuestDotStimuli/blob/wip/quest-pilot/Docs/"


def style_para(style, size, before=0, after=3, bold=None, italic=None, color=None, line=1.1):
    style.font.name = FONT
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    # East-Asian font slot too, or Word/Pages may fall back
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def style_table(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFC5CC")
        borders.append(el)
    tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 20), ("bottom", 20), ("left", 70), ("right", 70)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(w))
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    tblPr.append(margins)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(width)
    # Stretch the column grid to the text width, keeping pandoc's proportions
    grid = table._tbl.tblGrid
    cols = grid.findall(qn("w:gridCol")) if grid is not None else []
    total = sum(int(c.get(qn("w:w")) or 0) for c in cols)
    if total:
        text_width = int(7.1 * 1440)
        for c, row_cells in zip(cols, zip(*[r.cells for r in table.rows])):
            w = int(int(c.get(qn("w:w"))) * text_width / total)
            c.set(qn("w:w"), str(w))
            for cell in row_cells:
                cell.width = w * 635  # twips -> EMU
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            if r == 0:
                set_cell_shading(cell, "E8EDF2")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(9)
                    if r == 0:
                        run.font.bold = True


def main(src, dst):
    dst = Path(dst).expanduser()
    dst.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        # Relative links to other docs only work on GitHub; point them there
        md = re.sub(r"\]\(([\w./-]+\.md)\)", lambda m: f"]({REPO_DOCS}{m.group(1)})", Path(src).read_text())
        tmp_md = Path(tmp) / "src.md"
        tmp_md.write_text(md)
        raw = Path(tmp) / "raw.docx"
        subprocess.run(["pandoc", str(tmp_md), "-o", str(raw)], check=True)
        doc = Document(raw)

    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(0.7)
        section.top_margin = section.bottom_margin = Inches(0.6)

    styles = {s.name: s for s in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in styles:
            style_para(styles[name], 10, after=3 if name != "Compact" else 1)
    if "Title" in styles:
        style_para(styles["Title"], 20, after=2, bold=True, color=ACCENT)
    if "Subtitle" in styles:
        style_para(styles["Subtitle"], 10, after=8, italic=True, color=RGBColor(0x55, 0x5B, 0x63))
    if "Heading 1" in styles:
        style_para(styles["Heading 1"], 16, before=4, after=3, bold=True, color=ACCENT)
    if "Heading 2" in styles:
        style_para(styles["Heading 2"], 13, before=10, after=3, bold=True, color=ACCENT)
    if "Heading 3" in styles:
        style_para(styles["Heading 3"], 11, before=7, after=2, bold=True, color=ACCENT)
    if "Verbatim Char" in styles:
        v = styles["Verbatim Char"]
        v.font.name = MONO
        v.font.size = Pt(8.5)
        rpr = v.element.get_or_add_rPr()
        f = rpr.find(qn("w:rFonts"))
        if f is None:
            f = OxmlElement("w:rFonts")
            rpr.append(f)
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            f.set(qn(k), MONO)
    if "Source Code" in styles:
        s = styles["Source Code"]
        s.font.size = Pt(8.5)
        s.paragraph_format.space_after = Pt(4)

    # Drop pandoc's empty spacer paragraphs
    for p in list(doc.paragraphs):
        if not p.text.strip() and not p._p.xpath(".//w:drawing"):
            p._p.getparent().remove(p._p)

    for table in doc.tables:
        style_table(table)

    doc.save(dst)
    print(f"wrote {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
